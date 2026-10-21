# -*- coding: utf-8 -*-
"""Sell in 資料匯出(唯讀)— 沿用 NS 匯入程式的 REST / SuiteQL 連線(TBA、API User Role)。

只送 SuiteQL 查詢,不會建立、修改、刪除 NetSuite 任何資料。
口徑跟 NetSuite 報表 405「Sell in 成本分析」相同(2026/10/14 第十二版),詳見 SellIn匯出_說明.md。

用法(在 NS匯入程式 資料夾執行;預設連正式區):
  python sellin_export.py --probe                        先檢查 API 角色看得到哪些資料(建議第一次先跑)
  python sellin_export.py --from 2026-07 --to 2026-09    匯出 7~9 月
  python sellin_export.py --from 2026-09 --print-sql     只印出 SQL,不連線
選項:
  --to     不填＝跟 --from 同一個月
  --env    prod(預設)/ sandbox
  --out    輸出 xlsx 路徑(預設 data/sellin/SellIn_起_迄_時間.xlsx)
  --csv    另外輸出 CSV(UTF-8 BOM,給 Power BI 用)
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import os
import sys
import time

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

KIT_PTYPE_ID = "5"      # DR_存貨型態「組包」的內部 ID(API 角色讀不到這個自訂記錄,不能像 405 用名稱比對)
PRICE_LEVEL = 1         # 價格級別 1 = 建議售價
TWD = 1                 # 幣別 1 = 台幣(405 第十一版以前沒限定幣別,多幣別價格會讓組包用量灌水;第十二版已同步修正)
CHUNK = 900             # IN (...) 一批的筆數(Oracle 上限 1000)
LOAN_WORD = "領用"


# ------------------------------------------------------------------ SQL
def month_list(ym_from: str, ym_to: str) -> list[tuple[str, str, str]]:
    y, m = map(int, ym_from.split("-"))
    y2, m2 = map(int, ym_to.split("-"))
    out = []
    while (y, m) <= (y2, m2):
        first = dt.date(y, m, 1)
        nxt = dt.date(y + (1 if m == 12 else 0), 1 if m == 12 else m + 1, 1)
        out.append((f"{y:04d}-{m:02d}", first.isoformat(), (nxt - dt.timedelta(days=1)).isoformat()))
        y, m = nxt.year, nxt.month
    return out


def _between(col: str, d1: str, d2: str) -> str:
    return f"({col} >= TO_DATE('{d1}','YYYY-MM-DD') AND {col} <= TO_DATE('{d2}','YYYY-MM-DD'))"


def sellin_sql(d1: str, d2: str) -> str:
    """一個月的 Sell in(客戶 × 品項 × 是否贈備品),跟 405 costSql 同一套取數,只是不用 WITH。

    rev：發票／貸方通知單的行,金額取該行總帳(已過帳),數量拆銷貨(發票)／銷退(貸方通知單)。
    cst：發票／貸方通知單／出貨單的 COGS 分錄 ＋ 退貨授權項目收貨的 CUSTOMERRETURNVARIANCE(退貨入庫沖回)。
    paid：發票看同單同品項收入行金額是否全 0;出貨單看原銷售訂單單價;退貨收貨看退貨授權單價。
    """
    r = _between("t.trandate", d1, d2)
    return f"""
SELECT b.ent, b.itm, b.paid,
  SUM(b.amt) AS amt, SUM(b.cogs) AS cogs, SUM(b.sqty) AS sqty, SUM(b.rqty) AS rqty, SUM(b.cqty) AS cqty
FROM (
  SELECT t.entity AS ent, tl.item AS itm,
    CASE WHEN tl.netamount <> 0 THEN NVL(ral.credit, 0) - NVL(ral.debit, 0) ELSE 0 END AS amt,
    0 AS cogs,
    CASE WHEN t.type = 'CustInvc' AND tl.accountinglinetype IN ('INCOME','INTERCOINCOME')
         THEN -NVL(tl.quantity, 0) ELSE 0 END AS sqty,
    CASE WHEN t.type = 'CustCred' AND tl.accountinglinetype IN ('INCOME','INTERCOINCOME')
         THEN -NVL(tl.quantity, 0) ELSE 0 END AS rqty,
    0 AS cqty,
    CASE WHEN tl.netamount <> 0 THEN 1 ELSE 0 END AS paid
  FROM transactionline tl
  JOIN transaction t ON t.id = tl.transaction
  LEFT JOIN transactionaccountingline ral
    ON ral.transaction = tl.transaction AND ral.transactionline = tl.id AND ral.accountingbook = 1
  WHERE t.type IN ('CustInvc','CustCred') AND tl.mainline = 'F' AND tl.taxline = 'F'
    AND {r} AND ral.posting = 'T'
  UNION ALL
  SELECT t.entity AS ent, tl.item AS itm,
    0 AS amt,
    NVL(tal.debit, 0) - NVL(tal.credit, 0) AS cogs,
    0 AS sqty, 0 AS rqty,
    NVL(tl.quantity, 0) AS cqty,
    CASE WHEN t.type = 'ItemShip' THEN NVL(sp.paid, 0)
         WHEN t.type = 'ItemRcpt' THEN NVL(rp.paid, 1)
         ELSE NVL(ip.paid, 1) END AS paid
  FROM transactionline tl
  JOIN transaction t ON t.id = tl.transaction
  JOIN transactionaccountingline tal ON tal.transaction = tl.transaction AND tal.transactionline = tl.id
  JOIN account a ON a.id = tal.account
  LEFT JOIN (
    SELECT tl2.transaction AS trn, tl2.item AS itm,
      MAX(CASE WHEN tl2.netamount <> 0 THEN 1 ELSE 0 END) AS paid
    FROM transactionline tl2 JOIN transaction t2 ON t2.id = tl2.transaction
    WHERE t2.type IN ('CustInvc','CustCred') AND tl2.accountinglinetype IN ('INCOME','INTERCOINCOME')
      AND {_between("t2.trandate", d1, d2)}
    GROUP BY tl2.transaction, tl2.item
  ) ip ON ip.trn = tl.transaction AND ip.itm = tl.item
  LEFT JOIN (
    SELECT ln.nextdoc AS trn, so.item AS itm, MAX(CASE WHEN so.rate > 0 THEN 1 ELSE 0 END) AS paid
    FROM nexttransactionlinelink ln
    JOIN transaction pt ON pt.id = ln.previousdoc
    JOIN transactionline so ON so.transaction = ln.previousdoc AND so.id = ln.previousline
    WHERE pt.type = 'SalesOrd'
    GROUP BY ln.nextdoc, so.item
  ) sp ON sp.trn = tl.transaction AND sp.itm = tl.item
  LEFT JOIN (
    SELECT ln.nextdoc AS trn, rl.item AS itm, MAX(CASE WHEN rl.rate > 0 THEN 1 ELSE 0 END) AS paid
    FROM nexttransactionlinelink ln
    JOIN transaction pt ON pt.id = ln.previousdoc
    JOIN transactionline rl ON rl.transaction = ln.previousdoc AND rl.id = ln.previousline
    WHERE pt.type = 'RtnAuth'
    GROUP BY ln.nextdoc, rl.item
  ) rp ON rp.trn = tl.transaction AND rp.itm = tl.item
  WHERE a.accttype = 'COGS' AND tal.posting = 'T'
    AND ((tl.accountinglinetype = 'COGS' AND t.type IN ('CustInvc','CustCred','ItemShip'))
      OR (tl.accountinglinetype = 'CUSTOMERRETURNVARIANCE' AND t.type = 'ItemRcpt'))
    AND {r}
) b
JOIN item i ON i.id = b.itm
WHERE i.itemtype <> 'NonInvtPart'
GROUP BY b.ent, b.itm, b.paid
""".strip()


def items_sql(ids: list[str]) -> str:
    return ("SELECT i.id, i.itemid, i.custitem_dr_itemname1 AS code, i.custitem_dr_itemname2 AS spec,"
            " i.displayname, i.itemtype, i.isinactive,"
            " BUILTIN.DF(i.cseg2) AS series, BUILTIN.DF(i.custitem_dr_dosageform) AS dose,"
            " i.custitem_dr_dosageform AS dose_id,"
            " i.custitem_dr_product_type AS ptype_id, BUILTIN.DF(i.custitem_dr_product_type) AS ptype"
            f" FROM item i WHERE i.id IN ({','.join(ids)})")


def customers_sql(ids: list[str]) -> str:
    return ("SELECT c.id, c.entityid, c.companyname, BUILTIN.DF(c.category) AS category"
            f" FROM customer c WHERE c.id IN ({','.join(ids)})")


def bom_sql(kit_ids: list[str]) -> str:
    return ("SELECT kr.asm, bc.item AS comp, SUM(bc.quantity) AS qty"
            " FROM (SELECT ab.assembly AS asm, MAX(ab.currentrevision) AS rev FROM assemblyitembom ab"
            f"       WHERE ab.assembly IN ({','.join(kit_ids)}) GROUP BY ab.assembly) kr"
            " JOIN bomrevisioncomponent bc ON bc.bomrevision = kr.rev"
            " JOIN item ci ON ci.id = bc.item"
            " WHERE ci.custitem_dr_dosageform IS NOT NULL"
            " GROUP BY kr.asm, bc.item")


def price_sql(ids: list[str]) -> str:
    return ("SELECT pr.item, MAX(pr.unitprice) AS px FROM pricing pr"
            f" WHERE pr.pricelevel = {PRICE_LEVEL} AND pr.currency = {TWD}"  # ⚠ 別加 pr.quantity = 1:放在 WHERE 會一筆都查不到(SuiteQL 怪癖)
            f" AND pr.item IN ({','.join(ids)}) GROUP BY pr.item")


# ------------------------------------------------------------------ 小工具
def num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def sid(v) -> str:
    return "" if v is None else str(v)


def chunks(seq: list, n: int = CHUNK):
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def q(client, sql: str, what: str) -> list[dict]:
    ok, rows, err = client.suiteql(sql, limit=200000)
    if not ok:
        raise RuntimeError(f"{what} 查詢失敗:{err[:600]}")
    return [{k.lower(): v for k, v in r.items() if k != "links"} for r in rows]


# ------------------------------------------------------------------ 取數
def fetch(client, months: list[tuple[str, str, str]], log=print) -> dict:
    facts = []
    for ym, d1, d2 in months:
        t0 = time.time()
        rows = q(client, sellin_sql(d1, d2), f"{ym} Sell in")
        for r in rows:
            facts.append({
                "ym": ym, "ent": sid(r.get("ent")), "itm": sid(r.get("itm")), "paid": int(num(r.get("paid"))),
                "amt": num(r.get("amt")), "cogs": num(r.get("cogs")),
                "sqty": num(r.get("sqty")), "rqty": num(r.get("rqty")), "cqty": num(r.get("cqty")),
            })
        log(f"  {ym}:{len(rows):,} 列({time.time() - t0:.0f} 秒)")

    item_ids = sorted({f["itm"] for f in facts if f["itm"]}, key=int)
    items = {}
    for ch in chunks(item_ids):
        for r in q(client, items_sql(ch), "品項主檔"):
            items[sid(r["id"])] = r

    kit_ids = [i for i in item_ids if sid(items.get(i, {}).get("ptype_id")) == KIT_PTYPE_ID]
    bom = []
    for ch in chunks(kit_ids):
        bom += q(client, bom_sql(ch), "組包 BOM")
    comp_ids = sorted({sid(b["comp"]) for b in bom}, key=int)
    missing = [c for c in comp_ids if c not in items]
    for ch in chunks(missing):
        for r in q(client, items_sql(ch), "成分主檔"):
            items[sid(r["id"])] = r
    prices = {}
    for ch in chunks(comp_ids):
        for r in q(client, price_sql(ch), "建議售價"):
            prices[sid(r["item"])] = num(r.get("px"))

    cust_ids = sorted({f["ent"] for f in facts if f["ent"]}, key=int)
    customers = {}
    for ch in chunks(cust_ids):
        for r in q(client, customers_sql(ch), "客戶"):
            customers[sid(r["id"])] = r
    return {"facts": facts, "items": items, "bom": bom, "prices": prices, "customers": customers}


# ------------------------------------------------------------------ 拆組(同 405 explodeKits,價格限定台幣)
def kit_alloc(bom: list[dict], prices: dict) -> dict:
    by_kit: dict[str, list] = {}
    for b in bom:
        by_kit.setdefault(sid(b["asm"]), []).append(
            {"comp": sid(b["comp"]), "qty": num(b["qty"]), "px": prices.get(sid(b["comp"]), 0.0)})
    out = {}
    for kit, comps in by_kit.items():
        wsum = sum(c["qty"] * c["px"] for c in comps)
        qsum = sum(c["qty"] for c in comps)
        if wsum <= 0 and qsum <= 0:
            continue                       # 找不到可分攤的成分 → 組包原樣保留(同 405)
        out[kit] = [{"comp": c["comp"], "usage": c["qty"], "px": c["px"],
                     "share": (c["qty"] * c["px"] / wsum) if wsum > 0 else (c["qty"] / qsum)}
                    for c in comps]
    return out


def explode(facts: list[dict], alloc: dict) -> list[dict]:
    out = []
    for f in facts:
        a = alloc.get(f["itm"])
        if not a:
            g = dict(f)
            g["kit"] = ""
            out.append(g)
            continue
        for m in a:
            g = dict(f)
            g["kit"] = f["itm"]
            g["itm"] = m["comp"]
            for k in ("amt", "cogs"):
                g[k] = f[k] * m["share"]
            for k in ("sqty", "rqty", "cqty"):
                g[k] = f[k] * m["usage"]
            out.append(g)
    return out


# ------------------------------------------------------------------ 輸出
COLS = [
    ("年月", "ym"), ("客戶內部ID", "ent"), ("客戶代號", "cust_code"), ("客戶名稱", "cust_name"),
    ("通路別", "chancat"), ("領用", "loan"), ("品項內部ID", "itm"), ("品號", "code"), ("品名", "name"),
    ("產品代碼", "spec"), ("產品代碼前三碼", "p3"), ("產品代碼前六碼", "p6"), ("系列別", "series"),
    ("劑型分類", "dose"), ("存貨型態", "ptype"), ("贈備品(售價0)", "gift"),
    ("銷貨數量", "sqty"), ("銷退數量", "rqty"), ("銷貨淨量", "qty"),
    ("銷貨淨額", "amt"), ("銷貨成本", "cogs"), ("成本數量", "cqty"),
]
NUM_KEYS = {"sqty", "rqty", "qty", "amt", "cogs", "cqty"}


def enrich(rows: list[dict], data: dict, with_kit: bool) -> list[dict]:
    items, custs = data["items"], data["customers"]
    out = []
    for f in rows:
        it = items.get(f["itm"], {})
        cu = custs.get(f["ent"], {})
        spec = sid(it.get("spec"))
        name = sid(cu.get("companyname"))
        g = {
            "ym": f["ym"], "ent": f["ent"], "cust_code": sid(cu.get("entityid")), "cust_name": name,
            "chancat": sid(cu.get("category")), "loan": "Y" if LOAN_WORD in name else "",
            "itm": f["itm"], "code": sid(it.get("code")) or sid(it.get("itemid")),
            "name": sid(it.get("displayname")) or sid(it.get("itemid")),
            "spec": spec, "p3": spec[:3], "p6": spec[:6],
            "series": sid(it.get("series")) or "(未分類)", "dose": sid(it.get("dose")) or "(未分類)",
            "ptype": sid(it.get("ptype")), "gift": "" if f["paid"] == 1 else "Y",
            "sqty": f["sqty"], "rqty": f["rqty"], "qty": f["sqty"] + f["rqty"],
            "amt": f["amt"], "cogs": f["cogs"], "cqty": f["cqty"],
        }
        if with_kit:
            kit = items.get(f.get("kit", ""), {}) if f.get("kit") else {}
            g["kit_code"] = (sid(kit.get("code")) or sid(kit.get("itemid"))) if kit else ""
        out.append(g)
    out = [g for g in out if any(abs(g[k]) >= 0.005 for k in NUM_KEYS)]
    out.sort(key=lambda g: (g["ym"], g["cust_name"], g["code"], g["gift"]))
    return out


def totals(rows: list[dict]) -> dict:
    """預設篩選(排除贈備品、排除領用)的每月合計,方便對 405。"""
    tot: dict[str, dict] = {}
    for g in rows:
        if g["gift"] or g["loan"]:
            continue
        t = tot.setdefault(g["ym"], {k: 0.0 for k in NUM_KEYS})
        for k in NUM_KEYS:
            t[k] += g[k]
    return tot


def write_xlsx(path: str, detail: list[dict], exploded: list[dict], data: dict, alloc: dict, meta: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3D54")

    def sheet(ws, headers, rows, widths=None, numfmt=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) * 2 + 4)))
            fmt = (numfmt or {}).get(h)
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = fmt

    numfmt = {h: "#,##0" for h, k in COLS if k in NUM_KEYS}
    numfmt["分攤比例"] = "0.0000"
    numfmt["建議售價(台幣)"] = "#,##0.00"
    widths = {"客戶名稱": 22, "品名": 36, "產品代碼": 20, "來源組包品號": 18}

    ws = wb.active
    ws.title = "說明"
    for line in meta["lines"]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120

    sheet(wb.create_sheet("SellIn明細"), [h for h, _ in COLS],
          [[g[k] for _, k in COLS] for g in detail], widths, numfmt)
    hx = [h for h, _ in COLS] + ["來源組包品號"]
    sheet(wb.create_sheet("SellIn拆組"), hx,
          [[g[k] for _, k in COLS] + [g.get("kit_code", "")] for g in exploded], widths, numfmt)

    items = data["items"]
    brow = []
    for kit, comps in sorted(alloc.items(), key=lambda kv: sid(items.get(kv[0], {}).get("itemid"))):
        k = items.get(kit, {})
        for m in comps:
            c = items.get(m["comp"], {})
            brow.append([sid(k.get("code")) or sid(k.get("itemid")), sid(k.get("displayname")),
                         sid(c.get("code")) or sid(c.get("itemid")), sid(c.get("displayname")),
                         sid(c.get("series")), sid(c.get("dose")), m["usage"], m["px"], m["share"]])
    sheet(wb.create_sheet("組包BOM"), ["組包品號", "組包品名", "成分品號", "成分品名", "成分系列別", "成分劑型",
                                     "用量", "建議售價(台幣)", "分攤比例"], brow,
          {"組包品名": 36, "成分品名": 36}, {"用量": "#,##0.####", "建議售價(台幣)": "#,##0.00", "分攤比例": "0.0000"})

    irow = [[sid(r.get("id")), sid(r.get("itemid")), sid(r.get("code")), sid(r.get("displayname")),
             sid(r.get("spec")), sid(r.get("series")), sid(r.get("dose")), sid(r.get("ptype")),
             sid(r.get("itemtype")), "Y" if sid(r.get("isinactive")) == "T" else ""]
            for r in sorted(items.values(), key=lambda r: sid(r.get("itemid")))]
    sheet(wb.create_sheet("品項主檔"), ["內部ID", "項目名稱", "品號", "品名", "產品代碼", "系列別", "劑型分類",
                                     "存貨型態", "項目類型", "停用"], irow, {"品名": 36, "項目名稱": 40})
    crow = [[sid(r.get("id")), sid(r.get("entityid")), sid(r.get("companyname")), sid(r.get("category"))]
            for r in sorted(data["customers"].values(), key=lambda r: sid(r.get("companyname")))]
    sheet(wb.create_sheet("客戶"), ["內部ID", "客戶代號", "客戶名稱", "通路別"], crow, {"客戶名稱": 30})

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def write_csv(path: str, rows: list[dict], with_kit: bool):
    heads = [h for h, _ in COLS] + (["來源組包品號"] if with_kit else [])
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(heads)
        for g in rows:
            w.writerow([g[k] for _, k in COLS] + ([g.get("kit_code", "")] if with_kit else []))


# ------------------------------------------------------------------ 權限探測
# 管理員在 2026/10/14 查到的筆數(2026-09),只當參考;API 角色看到 0 筆通常就是缺權限(SuiteQL 不會報錯)
def probes(d1: str, d2: str) -> list[tuple[str, str, str, int]]:
    r = _between("t.trandate", d1, d2)
    line = "FROM transactionline tl JOIN transaction t ON t.id = tl.transaction WHERE "
    return [
        ("發票 收入行", "發票(檢視)",
         f"SELECT COUNT(*) AS n {line}t.type = 'CustInvc' AND tl.mainline = 'F' AND tl.taxline = 'F'"
         f" AND tl.accountinglinetype IN ('INCOME','INTERCOINCOME') AND {r}", 8332),
        ("貸方通知單 收入行(銷退)", "貸方通知單(檢視)",
         f"SELECT COUNT(*) AS n {line}t.type = 'CustCred' AND tl.mainline = 'F' AND tl.taxline = 'F'"
         f" AND tl.accountinglinetype IN ('INCOME','INTERCOINCOME') AND {r}", 1163),
        ("出貨單 成本行", "項目交貨(檢視)",
         f"SELECT COUNT(*) AS n {line}t.type = 'ItemShip' AND tl.accountinglinetype = 'COGS' AND {r}", 3552),
        ("退貨收貨 沖成本行", "項目收貨(檢視)",
         f"SELECT COUNT(*) AS n {line}t.type = 'ItemRcpt' AND tl.accountinglinetype = 'CUSTOMERRETURNVARIANCE' AND {r}",
         1025),
        ("發票 總帳分錄(已過帳)", "發票＋帳戶(檢視)",
         "SELECT COUNT(*) AS n FROM transactionaccountingline tal JOIN transaction t ON t.id = tal.transaction"
         f" WHERE t.type = 'CustInvc' AND tal.posting = 'T' AND tal.accountingbook = 1 AND {r}", 19092),
        ("銷售訂單→出貨 連結", "銷售訂單(檢視)",
         "SELECT COUNT(*) AS n FROM nexttransactionlinelink ln JOIN transaction pt ON pt.id = ln.previousdoc"
         f" JOIN transaction t ON t.id = ln.nextdoc WHERE pt.type = 'SalesOrd' AND t.type = 'ItemShip' AND {r}",
         3554),
        ("成本類科目", "帳戶(檢視)", "SELECT COUNT(*) AS n FROM account WHERE accttype = 'COGS'", 27),
        ("客戶", "客戶(檢視)", "SELECT COUNT(*) AS n FROM customer", 79),
        ("組包品項", "項目(檢視)", f"SELECT COUNT(*) AS n FROM item WHERE custitem_dr_product_type = {KIT_PTYPE_ID}",
         4513),
        ("BOM 成分", "用料表(檢視)", "SELECT COUNT(*) AS n FROM bomrevisioncomponent", 14795),
        ("台幣建議售價", "項目(檢視)",
         f"SELECT COUNT(*) AS n FROM pricing WHERE pricelevel = {PRICE_LEVEL} AND currency = {TWD}",
         5791),
        ("通路別名稱", "會計清單(檢視)", "SELECT COUNT(*) AS n FROM customercategory", 12),
        ("DR_產品系列", "自訂記錄 DR_產品系列", "SELECT COUNT(*) AS n FROM CUSTOMRECORD_CSEG2", 19),
        ("傳票(只有非品項收入／總帳表才需要)", "製作傳票(檢視)",
         f"SELECT COUNT(*) AS n FROM transaction t WHERE t.type = 'Journal' AND {r}", 199),
    ]


def run_probe(client, ym: str):
    (_, d1, d2), = month_list(ym, ym)
    print(f"== API 角色看得到的資料({ym},管理員 2026/10/14 的筆數當參考)")
    print(f"{'項目':<26}{'API':>10}{'管理員':>10}  判定")
    bad = []
    for label, perm, sql, base in probes(d1, d2):
        ok, rows, err = client.suiteql(sql, limit=1)
        if not ok:
            print(f"{label:<26}{'錯誤':>10}{base:>10,}  ✘ {err.splitlines()[0][:80]}")
            bad.append((label, perm))
            continue
        n = int(num(rows[0].get("n"))) if rows else 0
        mark = "✔" if n > 0 else "✘ 看不到 → 可能缺「" + perm + "」權限"
        if n == 0 and base > 0:
            bad.append((label, perm))
        print(f"{label:<26}{n:>10,}{base:>10,}  {mark}")
    print()
    print(f"== {ym} Sell in 合計(排除贈備品、排除領用、不拆組),跟 405 報表比")
    try:
        data = fetch(client, month_list(ym, ym), log=lambda s: None)
        rows = enrich(data["facts"], data, False)
        t = totals(rows).get(ym, {k: 0.0 for k in NUM_KEYS})
        exp = EXPECT_405.get(ym)
        for k, lab in (("amt", "銷貨淨額"), ("cogs", "銷貨成本"), ("sqty", "銷貨數量"), ("rqty", "銷退數量"), ("qty", "銷貨淨量")):
            ref = f"{exp[k]:>16,.0f}" if exp else f"{'-':>16}"
            same = ("✔" if exp and abs(round(t[k]) - exp[k]) <= 1 else ("✘" if exp else ""))
            print(f"{lab:<10}{t[k]:>16,.0f}{ref}  {same}")
        if exp:
            print("(右欄＝405 本期合計扣掉〔非品項收入〕與〔總帳成本差額〕＝商品行,2026/10/14 核對;月結還在動的月份會有落差)")
    except Exception as e:  # noqa: BLE001
        print("取數失敗:", e)
    if bad:
        print()
        print("需要請管理員在「API User Role」補的權限(檢視即可;沒補的話 SuiteQL 不會報錯,只會少資料):")
        for label, perm in bad:
            print(f"  - {perm}  ← {label}")


EXPECT_405 = {   # 405 本期(預設勾選),2026/10/14 用管理員跑同一支 SQL 核對過:三個月都一元不差
    "2026-07": {"amt": 110117707, "cogs": 21366112, "sqty": 221879, "rqty": -16199, "qty": 205680},
    "2026-08": {"amt": 177299615, "cogs": 32269699, "sqty": 288561, "rqty": -9055, "qty": 279506},
    "2026-09": {"amt": 117616663, "cogs": 22220933, "sqty": 223486, "rqty": -7614, "qty": 215872},
}


# ------------------------------------------------------------------ main
def main(argv=None):
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Sell in 資料匯出(唯讀,口徑同 405)")
    ap.add_argument("--from", dest="ym_from", help="起始年月 YYYY-MM")
    ap.add_argument("--to", dest="ym_to", help="結束年月 YYYY-MM(預設同起始)")
    ap.add_argument("--env", choices=["prod", "sandbox"], default="prod")
    ap.add_argument("--out", help="輸出 xlsx 路徑")
    ap.add_argument("--csv", action="store_true", help="另外輸出 CSV")
    ap.add_argument("--probe", action="store_true", help="檢查 API 角色看得到哪些資料")
    ap.add_argument("--probe-month", default="2026-09")
    ap.add_argument("--print-sql", action="store_true", help="只印 SQL")
    a = ap.parse_args(argv)

    if a.print_sql:
        ym = a.ym_from or a.probe_month
        for m, d1, d2 in month_list(ym, a.ym_to or ym):
            print(f"-- {m}\n{sellin_sql(d1, d2)}\n")
        return 0

    os.environ["NSIMP_ENV"] = "prod" if a.env == "prod" else "sandbox"
    from core.config import load_config
    from core.netsuite import NetSuiteClient
    cfg = load_config()
    client = NetSuiteClient(cfg)
    print(f"連線:{'正式區' if cfg['_env'] == 'prod' else 'sandbox'}(帳號 {cfg['netsuite']['account_id']}),只查詢不寫入")

    if a.probe:
        run_probe(client, a.probe_month)
        return 0

    if not a.ym_from:
        ap.error("請指定 --from YYYY-MM(或用 --probe)")
    months = month_list(a.ym_from, a.ym_to or a.ym_from)
    t0 = time.time()
    print(f"取數 {months[0][0]} ~ {months[-1][0]} …")
    data = fetch(client, months)
    alloc = kit_alloc(data["bom"], data["prices"])
    detail = enrich(data["facts"], data, False)
    exploded = enrich(explode(data["facts"], alloc), data, True)

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    out = a.out or os.path.join(BASE_DIR, "data", "sellin",
                                f"SellIn_{months[0][0]}_{months[-1][0]}_{stamp}.xlsx")
    t1, t2 = totals(detail), totals(exploded)
    lines = [
        f"Sell in 匯出 {months[0][0]} ~ {months[-1][0]}(產生時間 {dt.datetime.now():%Y-%m-%d %H:%M},"
        f"{'正式區' if cfg['_env'] == 'prod' else 'sandbox'})",
        "口徑同 NetSuite 報表 405「Sell in 成本分析」(第十二版),拆組的價格限定台幣。",
        "SellIn明細:客戶 × 品項 × 是否贈備品(不拆組)。SellIn拆組:組包依 BOM 拆成有劑型的成分(金額/成本依台幣建議售價×用量分攤,數量依用量)。",
        "銷貨數量＝發票數量;銷退數量＝貸方通知單數量(負值);銷貨淨量＝兩者相加。銷貨淨額取總帳(已過帳)。",
        "銷貨成本＝發票/貸方通知單/出貨單 COGS ＋ 退貨授權項目收貨的退貨入庫沖回;成本數量＝這些分錄行的數量。",
        "篩選欄:「贈備品(售價0)」=Y 為贈備品;「領用」=Y 為領用客戶。405 預設兩者都排除。",
        "不含:非品項收入(促銷費、折讓、權利金…)、〔總帳成本差額〕、總帳表。",
        "",
        "每月合計(排除贈備品、排除領用):",
        "年月 | 銷貨淨額 | 銷貨成本 | 銷貨數量 | 銷退數量 | 銷貨淨量 | 拆組後淨量",
    ]
    for ym, _, _ in months:
        x, y = t1.get(ym, {}), t2.get(ym, {})
        lines.append(f"{ym} | {x.get('amt', 0):,.0f} | {x.get('cogs', 0):,.0f} | {x.get('sqty', 0):,.0f} | "
                     f"{x.get('rqty', 0):,.0f} | {x.get('qty', 0):,.0f} | {y.get('qty', 0):,.0f}")
    write_xlsx(out, detail, exploded, data, alloc, {"lines": lines})
    if a.csv:
        base = os.path.splitext(out)[0]
        write_csv(base + "_明細.csv", detail, False)
        write_csv(base + "_拆組.csv", exploded, True)
    print("\n".join(lines[-(len(months) + 1):]))
    print(f"完成:{out}({time.time() - t0:.0f} 秒;明細 {len(detail):,} 列、拆組 {len(exploded):,} 列)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
